import os
import datetime
import openpyxl
from sqlalchemy.orm import Session
from database import SessionLocal, engine, Base
import models
from auth import get_password_hash

DATA_DIR = os.path.join(os.path.dirname(__file__), "data_migrasi")

def parse_date(val):
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.date() if isinstance(val, datetime.datetime) else val
    return None

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip().replace('\r\n', ' ').replace('\n', ' ')
    return s if s else None

def parse_float(val):
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0

def run_migration():
    print("Starting full database migration from Lark Base Excel files...")
    
    # 1. Reset database tables
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    
    db: Session = SessionLocal()
    seen_usernames = set()
    
    # 2. Seed Default Users
    default_users = [
        {"username": "superadmin", "password": "password", "role": "Superadmin"},
        {"username": "ir_user", "password": "password", "role": "IR"},
        {"username": "gov_user", "password": "password", "role": "Gov"},
        {"username": "pol_user", "password": "password", "role": "Pol"},
        {"username": "finance_user", "password": "password", "role": "Finance"},
        {"username": "viewer", "password": "password", "role": "Viewer"},
    ]
    for u in default_users:
        user_obj = models.User(
            username=u["username"],
            password_hash=get_password_hash(u["password"]),
            role=u["role"]
        )
        db.add(user_obj)
        seen_usernames.add(u["username"])
    db.commit()

    # 3. Seed Personil Internal as Users
    f_personil = os.path.join(DATA_DIR, "Project Management_Daftar Personil Internal_Daftar.xlsx")
    if os.path.exists(f_personil):
        wb_p = openpyxl.load_workbook(f_personil, data_only=True)
        ws_p = wb_p['Daftar Personil Internal']
        p_rows = list(ws_p.iter_rows(values_only=True))
        if p_rows:
            h_p = {str(name).strip(): i for i, name in enumerate(p_rows[0]) if name}
            for row in p_rows[1:]:
                nama = clean_str(row[h_p.get('Nama', 0)])
                divisi = clean_str(row[h_p.get('Divisi', 1)])
                jabatan = clean_str(row[h_p.get('Jabatan', 2)])
                if nama:
                    uname = nama.lower().replace(' ', '_').replace('.', '')
                    if uname not in seen_usernames:
                        role_map = {
                            "Public Policy": "Gov",
                            "Political Science": "Pol",
                            "Institutional Relationship": "IR",
                            "Finance & Administration": "Finance",
                            "Chief & Founder": "Superadmin",
                            "Data Science": "Viewer",
                            "System & Technology": "Viewer",
                            "People & Culture": "Viewer"
                        }
                        role = role_map.get(divisi, "Viewer") if divisi else "Viewer"
                        level = str(jabatan).strip().upper() if jabatan else "STAFF"
                        db.add(models.User(
                            username=uname,
                            password_hash=get_password_hash("password"),
                            role=role,
                            level=level,
                            divisi=divisi
                        ))
                        seen_usernames.add(uname)
            db.commit()

    # 4. Migrate Projects from "Project Management_Daftar Bidding_Kisi 11.xlsx"
    f_bidding = os.path.join(DATA_DIR, "Project Management_Daftar Bidding_Kisi 11.xlsx")
    wb_b = openpyxl.load_workbook(f_bidding, data_only=True)
    ws_b = wb_b['Daftar Bidding']
    b_rows = list(ws_b.iter_rows(values_only=True))
    h_b = {str(name).strip(): i for i, name in enumerate(b_rows[0]) if name}

    project_map = {} # nama_pekerjaan -> Project Model Instance

    for row in b_rows[1:]:
        nama_pekerjaan = clean_str(row[h_b.get('Nama Pekerjaan')])
        if not nama_pekerjaan:
            continue

        pemberi_kerja = clean_str(row[h_b.get('Pemberi Kerja')])
        satuan_kerja = clean_str(row[h_b.get('Satuan Kerja')])
        nilai_kontrak = parse_float(row[h_b.get('Nilai Kontrak')])
        
        # Clean Tahapan
        raw_tahapan = clean_str(row[h_b.get('Tahapan')]) or "Upload PQ"
        tahapan = raw_tahapan.split('.', 1)[-1].strip() if '.' in raw_tahapan else raw_tahapan

        # Clean Mekanisme
        mekanisme_val = clean_str(row[h_b.get('Mekanisme Penunjukan')])
        jenis_mekanisme = "PL" if mekanisme_val == "PL" else "Bidding"

        # Clean Status
        raw_status = clean_str(row[h_b.get('Status')])
        if jenis_mekanisme == "PL":
            status_project = "Menang"
        elif raw_status == "Menang":
            status_project = "Menang"
        elif raw_status == "Kalah":
            status_project = "Kalah"
        elif raw_status == "Batal":
            status_project = "Batal"
        elif raw_status in ["Tidak Memenuhi Nilai Ambang Batas", "Ambang Batas"]:
            status_project = "Tidak Memenuhi Ambang Batas"
        else:
            status_project = "Ongoing"

        lokasi = clean_str(row[h_b.get('Lokasi')]) or "Pusat"
        kategori_project = clean_str(row[h_b.get('Kategori Project')]) or "Gov"
        divisi_substansi = kategori_project

        deadline_pengumuman = parse_date(row[h_b.get('Deadline/Pengumuman')])
        keterangan_tender = clean_str(row[h_b.get('Keterangan Tender')])
        peringkat = clean_str(row[h_b.get('Peringkat')])
        
        pic = clean_str(row[h_b.get('PIC Project')]) or clean_str(row[h_b.get('PIC')])
        pic_ustek = clean_str(row[h_b.get('PIC Ustek')])
        tim_ustek = clean_str(row[h_b.get('Tim Penulis Ustek')])
        url_ustek = clean_str(row[h_b.get('URL Ustek')])

        status_penulisan_ustek = clean_str(row[h_b.get('Status Penulisan Ustek')])
        status_selesai_substansi = True if status_penulisan_ustek == "Selesai" else False
        deadline_penulisan_ustek = parse_date(row[h_b.get('Deadline Penulisan Ustek')])

        admin = clean_str(row[h_b.get('Admin')])
        tanggal_mulai_spk = parse_date(row[h_b.get('Tanggal Mulai SPK')])
        tanggal_spk_berakhir = parse_date(row[h_b.get('Tanggal SPK Berakhir')])

        link_spk = clean_str(row[h_b.get('Link SPK')]) or clean_str(row[h_b.get('SPK')])
        link_bast = clean_str(row[h_b.get('Link BAST')]) or clean_str(row[h_b.get('BAST')])
        link_referensi = clean_str(row[h_b.get('Link Referensi')]) or clean_str(row[h_b.get('Referensi')])

        project_obj = models.Project(
            nama_pekerjaan=nama_pekerjaan,
            pemberi_kerja=pemberi_kerja,
            satuan_kerja=satuan_kerja,
            nilai_kontrak=nilai_kontrak,
            jenis_mekanisme=jenis_mekanisme,
            tahapan=tahapan,
            status_project=status_project,
            lokasi=lokasi,
            kategori_project=kategori_project,
            divisi_substansi=divisi_substansi,
            deadline_pengumuman=deadline_pengumuman,
            keterangan_tender=keterangan_tender,
            peringkat=peringkat,
            pic=pic,
            pic_ustek=pic_ustek,
            tim_ustek=tim_ustek,
            url_ustek=url_ustek,
            status_penulisan_ustek=status_penulisan_ustek,
            status_selesai_substansi=status_selesai_substansi,
            deadline_penulisan_ustek=deadline_penulisan_ustek,
            admin=admin,
            tanggal_mulai_spk=tanggal_mulai_spk,
            tanggal_spk_berakhir=tanggal_spk_berakhir,
            link_spk=link_spk,
            link_bast=link_bast,
            link_referensi=link_referensi
        )

        db.add(project_obj)
        db.flush() # Populate project_obj.id
        project_map[nama_pekerjaan] = project_obj

        # Process "Task Project" column for unique project stages
        task_project_str = clean_str(row[h_b.get('Task Project')])
        if task_project_str:
            stages_list = [s.strip() for s in task_project_str.split(',') if s.strip()]
            for s_idx, stage_name in enumerate(stages_list):
                st_status = "Selesai" if status_selesai_substansi else ("Ongoing" if s_idx == 0 else "Pending")
                stage_obj = models.ProjectStage(
                    project_id=project_obj.id,
                    nama_tahapan=stage_name,
                    status=st_status
                )
                db.add(stage_obj)

        # Process "Manajemen Penagihan" for billings
        mgm_penagihan = clean_str(row[h_b.get('Manajemen Penagihan')])
        if nilai_kontrak > 0:
            billing_status = "Sudah dibayarkan" if admin == "Selesai" or link_bast else "Belum ditagih"
            b_name = mgm_penagihan if mgm_penagihan else "Termin 1 (Pelunasan)"
            db.add(models.Billing(
                project_id=project_obj.id,
                nama=b_name,
                nominal=nilai_kontrak,
                status=billing_status
            ))

    db.commit()
    print(f"Successfully migrated {len(project_map)} projects into database!")

    # 5. Migrate Tasks from "Project Management_Manajemen Project_Kisi 3.xlsx"
    f_tasks = os.path.join(DATA_DIR, "Project Management_Manajemen Project_Kisi 3.xlsx")
    if os.path.exists(f_tasks):
        wb_t = openpyxl.load_workbook(f_tasks, data_only=True)
        ws_t = wb_t['Manajemen Project']
        t_rows = list(ws_t.iter_rows(values_only=True))
        h_t = {str(name).strip(): i for i, name in enumerate(t_rows[0]) if name}

        tasks_added = 0
        for row in t_rows[1:]:
            task_name = clean_str(row[h_t.get('Task')])
            p_name = clean_str(row[h_t.get('Project')])
            if not task_name or not p_name:
                continue

            target_project = project_map.get(p_name)
            if target_project:
                status_t = clean_str(row[h_t.get('Status')]) or "Ongoing"
                deadline_t = parse_date(row[h_t.get('Deadline')])
                tgl_mulai_t = parse_date(row[h_t.get('Tanggal Mulai')])
                tgl_selesai_t = parse_date(row[h_t.get('Tanggal Selesai')])
                pic_t = clean_str(row[h_t.get('PIC')])
                jenis_tahapan = clean_str(row[h_t.get('Jenis Tahapan')])

                db.add(models.Task(
                    project_id=target_project.id,
                    task=task_name,
                    deskripsi=clean_str(row[h_t.get('Deskripsi')]),
                    jenis_tahapan=jenis_tahapan,
                    status=status_t,
                    deadline=deadline_t,
                    tanggal_mulai=tgl_mulai_t,
                    tanggal_selesai=tgl_selesai_t,
                    pic=pic_t
                ))
                tasks_added += 1

        db.commit()
        print(f"Successfully migrated {tasks_added} tasks!")
    
    db.close()

if __name__ == "__main__":
    run_migration()
