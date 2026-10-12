import os
from datetime import date
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy.orm import Session

import auth
import models
from database import get_db, SessionLocal

router = APIRouter(tags=["Tenaga Ahli"])

EDIT_ROLES = ["Gov", "Pol", "IR", "Systech"]
DATA_FILE = os.path.join(os.path.dirname(__file__), "data_ta", "database_tenaga_ahli.xlsx")


# ----------------- SCHEMAS -----------------

class PendidikanBase(BaseModel):
    jenjang: str
    jurusan: Optional[str] = None
    universitas: Optional[str] = None
    tahun_lulus: Optional[int] = None

class PendidikanResponse(PendidikanBase):
    id: int
    class Config:
        from_attributes = True

class SertifikasiBase(BaseModel):
    nama_sertifikasi: str
    tahun: Optional[int] = None

class SertifikasiResponse(SertifikasiBase):
    id: int
    class Config:
        from_attributes = True

class PenugasanBase(BaseModel):
    project_id: int
    posisi: Optional[str] = None
    fee: Optional[float] = None
    tanggal_mulai: Optional[date] = None
    tanggal_selesai: Optional[date] = None
    status: Optional[str] = "Diajukan"
    catatan: Optional[str] = None

class PenugasanUpdate(BaseModel):
    project_id: Optional[int] = None
    posisi: Optional[str] = None
    fee: Optional[float] = None
    tanggal_mulai: Optional[date] = None
    tanggal_selesai: Optional[date] = None
    status: Optional[str] = None
    catatan: Optional[str] = None

class PenugasanResponse(PenugasanBase):
    id: int
    ta_id: int
    nama_pekerjaan: Optional[str] = None
    nama_ta: Optional[str] = None
    kode_ta: Optional[str] = None

class TABase(BaseModel):
    nama: str
    kategori_bidang: Optional[str] = None
    status: Optional[str] = "Eksternal"
    no_hp: Optional[str] = None
    bidang_keahlian: Optional[str] = None
    domisili: Optional[str] = None
    pengalaman_tahun: Optional[float] = None
    kategori_pengalaman: Optional[str] = None
    prioritas: Optional[str] = None
    catatan: Optional[str] = None

class TACreate(TABase):
    pendidikan: List[PendidikanBase] = []
    sertifikasi: List[SertifikasiBase] = []

class TAUpdate(BaseModel):
    nama: Optional[str] = None
    kategori_bidang: Optional[str] = None
    status: Optional[str] = None
    no_hp: Optional[str] = None
    bidang_keahlian: Optional[str] = None
    domisili: Optional[str] = None
    pengalaman_tahun: Optional[float] = None
    kategori_pengalaman: Optional[str] = None
    prioritas: Optional[str] = None
    catatan: Optional[str] = None

class TAResponse(TABase):
    id: int
    kode: Optional[str] = None
    sumber_data: Optional[str] = None
    pendidikan: List[PendidikanResponse] = []
    sertifikasi: List[SertifikasiResponse] = []
    penugasan: List[PenugasanResponse] = []


def _penugasan_out(a: models.TAPenugasan) -> PenugasanResponse:
    return PenugasanResponse(
        id=a.id, ta_id=a.ta_id, project_id=a.project_id, posisi=a.posisi, fee=a.fee,
        tanggal_mulai=a.tanggal_mulai, tanggal_selesai=a.tanggal_selesai, status=a.status, catatan=a.catatan,
        nama_pekerjaan=a.project.nama_pekerjaan if a.project else None,
        nama_ta=a.tenaga_ahli.nama if a.tenaga_ahli else None,
        kode_ta=a.tenaga_ahli.kode if a.tenaga_ahli else None,
    )

def _ta_out(t: models.TenagaAhli) -> TAResponse:
    return TAResponse(
        id=t.id, kode=t.kode, nama=t.nama, kategori_bidang=t.kategori_bidang, status=t.status, no_hp=t.no_hp,
        bidang_keahlian=t.bidang_keahlian, domisili=t.domisili, pengalaman_tahun=t.pengalaman_tahun,
        kategori_pengalaman=t.kategori_pengalaman, prioritas=t.prioritas, sumber_data=t.sumber_data, catatan=t.catatan,
        pendidikan=[PendidikanResponse.model_validate(x) for x in sorted(t.pendidikan, key=lambda e: e.id)],
        sertifikasi=[SertifikasiResponse.model_validate(x) for x in sorted(t.sertifikasi, key=lambda e: e.id)],
        penugasan=[_penugasan_out(a) for a in sorted(t.penugasan, key=lambda e: e.id)],
    )

def _get_ta(db: Session, ta_id: int) -> models.TenagaAhli:
    t = db.query(models.TenagaAhli).filter(models.TenagaAhli.id == ta_id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Tenaga Ahli not found")
    return t


# ----------------- TENAGA AHLI -----------------

@router.get("/tenaga-ahli", response_model=List[TAResponse])
def list_tenaga_ahli(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    return [_ta_out(t) for t in db.query(models.TenagaAhli).order_by(models.TenagaAhli.id).all()]

@router.post("/tenaga-ahli", response_model=TAResponse)
def create_tenaga_ahli(payload: TACreate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    data = payload.model_dump(exclude={"pendidikan", "sertifikasi"})
    t = models.TenagaAhli(**data, sumber_data="Input manual")
    t.pendidikan = [models.TAPendidikan(**e.model_dump()) for e in payload.pendidikan]
    t.sertifikasi = [models.TASertifikasi(**e.model_dump()) for e in payload.sertifikasi]
    db.add(t)
    db.flush()
    t.kode = "TA-%04d" % t.id
    db.commit()
    db.refresh(t)
    return _ta_out(t)

@router.put("/tenaga-ahli/{ta_id}", response_model=TAResponse)
def update_tenaga_ahli(ta_id: int, payload: TAUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    t = _get_ta(db, ta_id)
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(t, k, v)
    db.commit()
    db.refresh(t)
    return _ta_out(t)

@router.delete("/tenaga-ahli/{ta_id}")
def delete_tenaga_ahli(ta_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    db.delete(_get_ta(db, ta_id))
    db.commit()
    return {"message": "Deleted"}


# ----------------- PENDIDIKAN & SERTIFIKASI (one-to-many) -----------------

def _get_child(db: Session, model, cid: int):
    e = db.query(model).filter(model.id == cid).first()
    if not e:
        raise HTTPException(status_code=404, detail="Data not found")
    return e

@router.post("/tenaga-ahli/{ta_id}/pendidikan", response_model=PendidikanResponse)
def add_pendidikan(ta_id: int, payload: PendidikanBase, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    _get_ta(db, ta_id)
    e = models.TAPendidikan(ta_id=ta_id, **payload.model_dump())
    db.add(e); db.commit(); db.refresh(e)
    return e

@router.put("/ta-pendidikan/{pid}", response_model=PendidikanResponse)
def update_pendidikan(pid: int, payload: PendidikanBase, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    e = _get_child(db, models.TAPendidikan, pid)
    for k, v in payload.model_dump().items():
        setattr(e, k, v)
    db.commit(); db.refresh(e)
    return e

@router.delete("/ta-pendidikan/{pid}")
def delete_pendidikan(pid: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    db.delete(_get_child(db, models.TAPendidikan, pid)); db.commit()
    return {"message": "Deleted"}

@router.post("/tenaga-ahli/{ta_id}/sertifikasi", response_model=SertifikasiResponse)
def add_sertifikasi(ta_id: int, payload: SertifikasiBase, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    _get_ta(db, ta_id)
    e = models.TASertifikasi(ta_id=ta_id, **payload.model_dump())
    db.add(e); db.commit(); db.refresh(e)
    return e

@router.put("/ta-sertifikasi/{sid}", response_model=SertifikasiResponse)
def update_sertifikasi(sid: int, payload: SertifikasiBase, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    e = _get_child(db, models.TASertifikasi, sid)
    for k, v in payload.model_dump().items():
        setattr(e, k, v)
    db.commit(); db.refresh(e)
    return e

@router.delete("/ta-sertifikasi/{sid}")
def delete_sertifikasi(sid: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    db.delete(_get_child(db, models.TASertifikasi, sid)); db.commit()
    return {"message": "Deleted"}


# ----------------- PENUGASAN (assign TA ke project) -----------------

@router.get("/ta-penugasan", response_model=List[PenugasanResponse])
def list_penugasan(db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    return [_penugasan_out(a) for a in db.query(models.TAPenugasan).order_by(models.TAPenugasan.id).all()]

@router.get("/projects/{project_id}/tenaga-ahli", response_model=List[PenugasanResponse])
def list_project_ta(project_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.get_current_user)):
    rows = db.query(models.TAPenugasan).filter(models.TAPenugasan.project_id == project_id).order_by(models.TAPenugasan.id).all()
    return [_penugasan_out(a) for a in rows]

@router.post("/tenaga-ahli/{ta_id}/penugasan", response_model=PenugasanResponse)
def assign_ta(ta_id: int, payload: PenugasanBase, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    _get_ta(db, ta_id)
    if not db.query(models.Project).filter(models.Project.id == payload.project_id).first():
        raise HTTPException(status_code=404, detail="Project not found")
    dup = db.query(models.TAPenugasan).filter(
        models.TAPenugasan.ta_id == ta_id, models.TAPenugasan.project_id == payload.project_id,
        models.TAPenugasan.status != "Batal").first()
    if dup:
        raise HTTPException(status_code=409, detail="TA sudah ditugaskan di project ini")
    a = models.TAPenugasan(ta_id=ta_id, **payload.model_dump())
    db.add(a); db.commit(); db.refresh(a)
    return _penugasan_out(a)

@router.put("/ta-penugasan/{aid}", response_model=PenugasanResponse)
def update_penugasan(aid: int, payload: PenugasanUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    a = _get_child(db, models.TAPenugasan, aid)
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(a, k, v)
    db.commit(); db.refresh(a)
    return _penugasan_out(a)

@router.delete("/ta-penugasan/{aid}")
def delete_penugasan(aid: int, db: Session = Depends(get_db), current_user: models.User = Depends(auth.check_role(EDIT_ROLES))):
    db.delete(_get_child(db, models.TAPenugasan, aid)); db.commit()
    return {"message": "Deleted"}


# ----------------- IMPORT AWAL DARI EXCEL -----------------

def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v

def _int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None

def _rows(wb, sheet):
    it = wb[sheet].iter_rows(values_only=True)
    head = next(it)
    for r in it:
        if any(c is not None for c in r):
            yield dict(zip(head, r))

def import_from_excel(db: Session, path: str = DATA_FILE) -> int:
    """Isi tabel tenaga ahli dari Excel gabungan. Hanya dijalankan bila tabel masih kosong."""
    if not os.path.exists(path) or db.query(models.TenagaAhli).count() > 0:
        return 0
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    by_kode = {}
    for r in _rows(wb, "Database TA"):
        kode = _clean(r.get("ID TA"))
        if not kode or not _clean(r.get("Nama")):
            continue
        t = models.TenagaAhli(
            kode=kode, nama=_clean(r.get("Nama")), kategori_bidang=_clean(r.get("Kategori (Bidang)")),
            status=_clean(r.get("Status")) or "Eksternal", no_hp=_clean(r.get("Nomor Handphone")),
            bidang_keahlian=_clean(r.get("Bidang Keahlian")), domisili=_clean(r.get("Domisili")),
            pengalaman_tahun=_clean(r.get("Pengalaman (Tahun)")),
            kategori_pengalaman=_clean(r.get("Kategori Pengalaman (Tahun)")),
            prioritas=_clean(r.get("Prioritas")), sumber_data=_clean(r.get("Sumber Data")), catatan=_clean(r.get("Catatan")),
        )
        db.add(t)
        by_kode[kode] = t
    db.flush()
    for r in _rows(wb, "Pendidikan"):
        t = by_kode.get(_clean(r.get("ID TA")))
        if t:
            db.add(models.TAPendidikan(ta_id=t.id, jenjang=_clean(r.get("Jenjang")) or "-", jurusan=_clean(r.get("Jurusan")),
                                       universitas=_clean(r.get("Universitas")), tahun_lulus=_int(r.get("Tahun Lulus"))))
    for r in _rows(wb, "Sertifikasi"):
        t = by_kode.get(_clean(r.get("ID TA")))
        if t and _clean(r.get("Nama Sertifikasi")):
            db.add(models.TASertifikasi(ta_id=t.id, nama_sertifikasi=_clean(r.get("Nama Sertifikasi")), tahun=_int(r.get("Tahun"))))
    db.commit()
    return len(by_kode)

def seed_if_empty():
    db = SessionLocal()
    try:
        n = import_from_excel(db)
        if n:
            print(f"[tenaga_ahli] imported {n} tenaga ahli")
    finally:
        db.close()
